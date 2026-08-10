"""修复 test_cases.xlsx：正确列索引更新 F-02 相关用例。

表头: 0=case_name 1=endpoint 2=method 3=body 4=expected_status 5=expected_code 6=expected_msg
"""
import openpyxl

wb = openpyxl.load_workbook("data/test_cases.xlsx")
ws = wb["Sheet1"]

headers = [c.value for c in ws[1]]
print("表头:", headers)

# 先打印全部行当前状态
for row in ws.iter_rows(min_row=2):
    print([c.value for c in row])

print("--- 修复 ---")
for row in ws.iter_rows(min_row=2):
    name = row[0].value
    if name == "用户不存在":
        row[3].value = '{"username":"not_exist","password":"123456"}'  # body
        row[4].value = 401   # expected_status
        row[5].value = -1    # expected_code
        row[6].value = "用户名或密码错误"  # expected_msg
        print("已修复: 用户不存在")
    elif name == "错错错":
        row[3].value = '{"username":"错错错","password":"错错错"}'  # body
        row[4].value = 401   # expected_status
        row[5].value = -1    # expected_code
        row[6].value = "用户名或密码错误"  # expected_msg
        print("已修复: 错错错")

wb.save("data/test_cases.xlsx")
print("--- 保存后验证 ---")
ws2 = wb["Sheet1"]
for row in ws2.iter_rows(min_row=2):
    print([c.value for c in row])
