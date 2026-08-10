"""更新 test_cases.xlsx：F-02 修复后，用户不存在/错错错 改为 401 + 用户名或密码错误。"""
import openpyxl

wb = openpyxl.load_workbook("data/test_cases.xlsx")
ws = wb["Sheet1"]

headers = [c.value for c in ws[1]]
print("表头:", headers)

updated = []
for row in ws.iter_rows(min_row=2):
    name = row[0].value
    if name in ("用户不存在", "错错错"):
        row[3].value = 401  # expected_status
        row[5].value = "用户名或密码错误"  # expected_msg
        updated.append(name)

wb.save("data/test_cases.xlsx")
print("已更新:", updated)
print("已保存 data/test_cases.xlsx")
